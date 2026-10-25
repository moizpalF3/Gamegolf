"""
Reads test_data.xlsx and turns it into overrides for test_config.json.

Sheets (all optional; header row is row 1, extra rows/blank rows are ignored):
  Credentials  columns: type | email | password      (type = valid / invalid)
  Signup       columns: setting | value              (first_name, last_name,
                                                      email_domain, password,
                                                      existing_email)
  Clubs        columns: club_name | tag_id

Play Golf scenarios (read by load_play_golf_scenarios):
  Courses      columns: course | hole | location | lat | lon
               The coordinate library: every named spot on every hole
               (Tee, Fairway, Fairway - missed left, Green, ...).
  Scenarios    columns: scenario_id | description | course | run
               One row per Play Golf test case; run = Yes / No.
  Shots        columns: scenario_id | hole | shot | location | club
               The shots to play: move the Medallion to `location` on
               `hole`, then swing `club`. Leave club blank to walk through a
               location without a shot. Holes are played in the order they
               first appear; shots within a hole in `shot` order.
"""

from pathlib import Path

from openpyxl import load_workbook


def _rows(ws):
    """Yield each data row as {lowercased header: stripped string}, skipping blank rows."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return
    keys = [str(h).strip().lower() if h is not None else "" for h in header]
    for row in rows:
        record = {k: ("" if v is None else str(v).strip()) for k, v in zip(keys, row) if k}
        if any(record.values()):
            yield record


def load_overrides(path):
    """Returns a partial config dict, e.g. {"login": {...}, "signup": {...}, "play_golf": {"clubs": [...]}}."""
    wb = load_workbook(Path(path), data_only=True, read_only=True)
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}
    overrides = {}

    if "credentials" in sheets:
        login = {}
        for r in _rows(sheets["credentials"]):
            kind = r.get("type", "").lower()
            if kind in ("valid", "invalid") and r.get("email"):
                login[f"{kind}_credentials"] = {"email": r["email"], "password": r.get("password", "")}
        if login:
            overrides["login"] = login

    if "signup" in sheets:
        signup = {r["setting"]: r.get("value", "") for r in _rows(sheets["signup"]) if r.get("setting")}
        signup = {k: v for k, v in signup.items() if v != ""}
        if signup:
            overrides["signup"] = signup

    if "clubs" in sheets:
        clubs = [
            {"name": r["club_name"], "tag_id": r["tag_id"]}
            for r in _rows(sheets["clubs"])
            if r.get("club_name") and r.get("tag_id")
        ]
        if clubs:
            overrides["play_golf"] = {"clubs": clubs}

    wb.close()
    return overrides


def _number(value, what):
    try:
        return float(value)
    except ValueError:
        raise ValueError(f"{what}: '{value}' isn't a number") from None


def load_play_golf_scenarios(path):
    """
    Builds one Play Golf scenario per enabled row of the Scenarios sheet:
      {"id", "description", "course", "holes_config": {"waypoints", "holes_info"}}
    where holes_config is in the same shape as a waypoints JSON file, so the
    Medallion simulator can play it as is. Returns [] if the workbook has no
    Scenarios sheet.
    """
    wb = load_workbook(Path(path), data_only=True, read_only=True)
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}
    if "scenarios" not in sheets:
        wb.close()
        return []
    for required in ("courses", "shots"):
        if required not in sheets:
            wb.close()
            raise ValueError(f"{path}: has a Scenarios sheet but no '{required.title()}' sheet")

    locations = {}
    for r in _rows(sheets["courses"]):
        key = (r.get("course", "").lower(), int(_number(r.get("hole", ""), "Courses: hole")),
               r.get("location", "").lower())
        locations[key] = (_number(r.get("lat", ""), f"Courses: lat for {key}"),
                          _number(r.get("lon", ""), f"Courses: lon for {key}"))

    shots_by_scenario = {}
    for r in _rows(sheets["shots"]):
        shots_by_scenario.setdefault(r.get("scenario_id", ""), []).append(r)

    scenarios = []
    for r in _rows(sheets["scenarios"]):
        scenario_id = r.get("scenario_id", "")
        if not scenario_id or r.get("run", "yes").lower() not in ("yes", "y", "true", "1"):
            continue
        course = r.get("course", "")
        rows = shots_by_scenario.get(scenario_id)
        if not rows:
            raise ValueError(f"Scenario {scenario_id} has no rows in the Shots sheet")

        # Holes in first-seen order; shots within a hole by their shot number.
        hole_order = []
        for s in rows:
            hole = int(_number(s.get("hole", ""), f"Shots ({scenario_id}): hole"))
            s["_hole"] = hole
            if hole not in hole_order:
                hole_order.append(hole)

        waypoints = []
        for hole in hole_order:
            hole_rows = [s for s in rows if s["_hole"] == hole]
            hole_rows.sort(key=lambda s: _number(s.get("shot") or "0", f"Shots ({scenario_id}): shot"))
            for s in hole_rows:
                location = s.get("location", "")
                coords = locations.get((course.lower(), hole, location.lower()))
                if coords is None:
                    known = sorted(loc for (c, h, loc) in locations if c == course.lower() and h == hole)
                    raise ValueError(
                        f"Scenario {scenario_id}: location '{location}' on hole {hole} of "
                        f"'{course}' isn't in the Courses sheet. Known locations: {known or 'none'}"
                    )
                wp = {"hole": hole, "label": location, "lat": coords[0], "lon": coords[1]}
                if s.get("club"):
                    wp["club"] = s["club"]
                waypoints.append(wp)

        scenarios.append({
            "id": scenario_id,
            "description": r.get("description", ""),
            "course": course,
            "holes_config": {
                "waypoints": waypoints,
                "holes_info": {str(h): {"course": course, "name": f"Hole {h}"} for h in hole_order},
            },
        })

    wb.close()
    return scenarios


def merge_config(config, overrides):
    """Overlays overrides onto config one level deep (so unrelated keys in a section survive)."""
    for section, values in overrides.items():
        if isinstance(values, dict) and isinstance(config.get(section), dict):
            config[section].update(values)
        else:
            config[section] = values
    return config


def apply_club_tags(waypoints, clubs):
    """
    Fills in "tag_id" on every waypoint that names a "club" instead, using the
    Clubs list ({"name", "tag_id"}). Waypoints that already have a tag_id, or
    no club at all (movement-only points), are left alone.
    """
    by_name = {c["name"].strip().lower(): c["tag_id"] for c in clubs}
    for wp in waypoints:
        club = wp.get("club")
        if club and "tag_id" not in wp:
            tag = by_name.get(club.strip().lower())
            if tag is None:
                raise ValueError(
                    f"Waypoint '{wp.get('label', '?')}' (hole {wp.get('hole', 1)}) uses club "
                    f"'{club}', which isn't in the Clubs sheet. Available: {sorted(c['name'] for c in clubs)}"
                )
            wp["tag_id"] = tag
    return waypoints
